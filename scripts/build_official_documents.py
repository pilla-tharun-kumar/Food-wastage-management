"""Build the three Milestone 1 briefing documents."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
GREEN = RGBColor(10, 105, 82)
INK = RGBColor(30, 49, 55)
MUTED = RGBColor(91, 112, 118)


def shade(cell, fill):
    elt = OxmlElement("w:shd")
    elt.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(elt)


def cell_text(cell, text, bold=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    par = cell.paragraphs[0]
    par.paragraph_format.space_before = Pt(3)
    par.paragraph_format.space_after = Pt(3)
    run = par.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    run.font.color.rgb = GREEN if bold else INK


def table(doc, headers, rows):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = True
    for cell, header in zip(tbl.rows[0].cells, headers):
        shade(cell, "EAF5EF")
        cell_text(cell, header, True)
    for index, values in enumerate(rows):
        for cell, value in zip(tbl.add_row().cells, values):
            if index % 2:
                shade(cell, "F3F7F8")
            cell_text(cell, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def document(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.72)
    sec.bottom_margin = Inches(.68)
    sec.left_margin = sec.right_margin = Inches(.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.16
    for name, size in (("Title", 23), ("Heading 1", 13), ("Heading 2", 10.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0) if name == "Title" else GREEN
        style.paragraph_format.space_before = Pt(13 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(9.5)
        style.paragraph_format.space_after = Pt(3)
    header = sec.header.paragraphs[0]
    header.text = "FOODBRIDGE  |  INVENTORY MANAGEMENT AND EXPIRY TRACKING"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = GREEN
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Milestone 1  |  14 October 2026  |  Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_paragraph(title, "Title")
    lead = doc.add_paragraph(subtitle)
    lead.paragraph_format.space_after = Pt(13)
    for run in lead.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = MUTED
    return doc


def h(doc, text):
    doc.add_paragraph(text, "Heading 1")


def p(doc, text):
    doc.add_paragraph(text)


def bullet(doc, text):
    doc.add_paragraph(text, "List Bullet")


def idea():
    doc = document("Project Problem Statement and Vision", "FoodBridge  |  AI Powered Food Waste Management Platform")
    p(doc, "FoodBridge addresses avoidable food loss in food businesses by making stock, expiry and surplus visible early enough for staff to act. The long term vision connects operational inventory data with demand prediction and redistribution. The present implementation establishes the inventory and expiry foundation for that vision.")
    h(doc, "The problem")
    p(doc, "Food service businesses and supermarkets often manage perishable stock across spreadsheets, manual counts and separate sales systems. Without a reliable batch level record, teams can miss expiry dates, overpurchase and identify surplus too late for safe redistribution. Recipient organizations face a separate discovery and coordination problem when available food is not listed in time.")
    h(doc, "Who the platform serves")
    table(doc, ["User group", "Need", "Planned value"], [
        ("Food businesses", "Know exact stock, expiry and movement", "Earlier use, transfer or donation decisions"),
        ("NGOs and community kitchens", "Discover usable surplus and arrange collection", "Timely access to available food"),
        ("Operations teams", "Track actions and measure results", "Auditable decisions and improvement over time"),
    ])
    h(doc, "Proposed solution")
    p(doc, "The full platform is planned as a two sided system. Businesses record inventory, receive expiry and future waste risk signals, and list suitable surplus. Recipient organizations find relevant donations and coordinate pickup. Analytics later connect stock decisions with recovery and waste reduction measures.")
    h(doc, "Why the first milestone matters")
    p(doc, "The first two weeks focus on trustworthy inventory data: products, batches, quantities, expiry dates, categories, stock history and donation records. These records make later forecasting and marketplace features technically possible, while providing immediate value through alerts and a usable dashboard.")
    doc.add_page_break()
    h(doc, "Eight week delivery path")
    table(doc, ["Period", "Focus", "Expected output"], [
        ("Weeks 1–2", "Inventory and expiry", "Manual, CSV and POS ready input; alerts; taxonomy; scanning; tenant records"),
        ("Weeks 3–4", "Waste prediction", "Historical data pipeline, demand forecasts, risk score and reorder advice"),
        ("Weeks 5–6", "Redistribution", "Surplus listings, matching, pickup coordination, business and NGO dashboards"),
        ("Weeks 7–8", "Finalization", "Integration, load and browser checks, documentation and final presentation"),
    ])
    h(doc, "Intended outcomes and evaluation")
    p(doc, "The intended outcome is earlier identification of at risk stock and faster transfer of usable surplus. In later milestones, the team can evaluate this through measures such as batches flagged before expiry, quantities donated, avoidable waste trends, forecast error and pickup completion. These are proposed measures; measured waste reduction is not yet available from the Week 1–2 prototype.")
    h(doc, "Current boundary")
    p(doc, "The working Milestone 1 application provides inventory management, expiry alerts, categorization, barcode and QR lookup, transaction history and completed donation records. AI forecasting, NGO access, matching, pickup scheduling and impact calculations belong to later phases and should not be presented as live features today.")
    h(doc, "Presentation message")
    p(doc, "FoodBridge begins with accurate stock and expiry records. Those records let a business act on food at risk today and create the data foundation for prediction and redistribution in the remaining milestones.")
    doc.save(OUT / "01_Project_Problem_Statement_and_Vision.docx")


def mvp():
    doc = document("Milestone 1 Minimum Viable Product", "Weeks 1 and 2  |  Scope, user journeys and acceptance criteria")
    p(doc, "The Week 1–2 MVP is a working, business scoped inventory and expiry application. It lets staff record stock, find batches nearing expiry, update quantities and retain an auditable history. It is suitable for a local demonstration and for integration work on later modules.")
    h(doc, "Delivered scope")
    table(doc, ["Capability", "What a user can do"], [
        ("Inventory", "Create products and batches; edit, search, filter and delete batches; record quantities, dates and locations"),
        ("Bulk and POS input", "Import a validated CSV file; use authenticated REST routes for stock changes from a POS client"),
        ("Expiry action", "Apply a business or batch alert threshold; see in app alerts and acknowledge them"),
        ("Food taxonomy", "Use seeded categories with perishability and storage guidance"),
        ("Scanning", "Look up a saved product by barcode or QR using camera, image or typed code"),
        ("Account separation", "Keep each business inventory, alerts, transactions and donations separate"),
        ("Traceability", "Review stock changes and log completed donations with automatic stock deduction"),
    ])
    h(doc, "Primary user journey")
    p(doc, "A business member signs in, creates a product and its first stock batch, then views its status on the dashboard. As the expiry date approaches, the batch appears in the expiring filter and in app alerts. Staff can adjust stock with a reason, identify it by scanning a code, or record a completed donation. The resulting stock movement remains in the activity history.")
    h(doc, "What makes this a viable first release")
    p(doc, "The application demonstrates an end to end operational loop: intake, classification, expiry visibility, action and recordkeeping. Browser forms and REST endpoints use the same business rules, so future clients can build on the existing API.")
    doc.add_page_break()
    h(doc, "Acceptance criteria and evidence")
    table(doc, ["Check", "Expected behavior", "Evidence"], [
        ("Stock input", "Valid records save; invalid dates, quantities and duplicates are rejected", "API and form tests"),
        ("CSV import", "One invalid row cancels the entire file with a row error", "Atomic rollback tests"),
        ("Expiry", "Today and threshold boundaries classify correctly; alerts refresh", "Date and alert lifecycle tests"),
        ("Business isolation", "Another business cannot read or change records", "Cross business tests"),
        ("History", "Adjustments and donations retain an audit snapshot", "Stock and donation tests"),
        ("Responsive UI", "Dashboard and navigation work at phone width", "Local browser verification"),
    ])
    h(doc, "API handoff")
    p(doc, "The application exposes internal Django REST Framework endpoints. Core routes include GET and POST /api/inventory/, GET and PUT or PATCH or DELETE /api/inventory/{id}/, GET /api/inventory/expiring/, POST /api/inventory/import-csv/, GET /api/alerts/, GET /api/transactions/ and GET or POST /api/donations/. POS clients use token authentication and a retry reference for stock adjustments.")
    h(doc, "MVP limits")
    bullet(doc, "Alerts are shown in the dashboard. Email, SMS and push delivery are not part of this release.")
    bullet(doc, "The POS contract is ready, but a vendor specific connector awaits POS details and credentials.")
    bullet(doc, "Phone camera behavior still needs a check on the target device over HTTPS; image and typed lookup work in the local demo.")
    bullet(doc, "Forecasting, risk scoring, marketplace listings, NGO matching and pickup scheduling are planned after Milestone 1.")
    h(doc, "Completion evidence")
    p(doc, "The local automated suite reports 29 passing tests. A browser walkthrough has checked sign in, stock forms, filtering, barcode and QR lookup, stock adjustment, responsive navigation and dashboard rendering. The application remains a local milestone build rather than a public production deployment.")
    doc.save(OUT / "02_Milestone_1_MVP.docx")


def status():
    doc = document("Project Status and Technical Walkthrough", "FoodBridge  |  Milestone 1 status as of 14 October 2026")
    p(doc, "Milestone 1 is implemented and verified locally. The project has a responsive dashboard, authenticated REST APIs, a tenant scoped database, scheduled expiry checks and operational history. The team can use the current build for a live review of Weeks 1–2 while keeping later AI and redistribution modules on the roadmap.")
    h(doc, "Current status")
    table(doc, ["Area", "Status", "Evidence or next step"], [
        ("Inventory and expiry", "Implemented locally", "Manual and CSV intake, API, alerts, taxonomy, scanner, history"),
        ("User interface", "Implemented locally", "Desktop and mobile layouts; collapsible sidebar; live charts"),
        ("Quality checks", "Passed locally", "29 automated tests; Django and JavaScript checks"),
        ("POS vendor connection", "Contract ready", "A named vendor adapter and credentials are still required"),
        ("Celery and Redis", "Configured, not live tested", "Local watcher is the default running scheduler"),
        ("Deployment and sign off", "Pending", "Repository push, target environment and mentor review remain"),
    ])
    h(doc, "Technology stack in use")
    table(doc, ["Layer", "Implementation", "Role"], [
        ("Interface", "Django templates, JavaScript, CSS", "Dashboard, forms, charts and mobile navigation"),
        ("API", "Django 5.2.18 and DRF 3.17.2", "Validation, authentication and inventory operations"),
        ("Database", "SQLite for local development", "Products, batches, alerts, history and donation logs"),
        ("Expiry scheduling", "Django command and local watcher", "Checks dates and refreshes in app alerts"),
        ("Optional scheduler", "Celery 5.6.3 with Redis 7.4.1", "Broker backed scheduled checks when services are deployed"),
        ("Scanning", "Bundled html5 qrcode library", "Camera, image and exact code lookup"),
    ])
    doc.add_page_break()
    h(doc, "How the application works")
    p(doc, "A signed in browser uses a same origin session and CSRF protection to call DRF endpoints. An external POS client can use a scoped token. The API validates requests, derives the business from the account membership, and calls transactional inventory services. Those services write inventory records and stock history to the database. Expiry checks update active alerts; the dashboard reads summary counts and refreshes them while visible.")
    h(doc, "Data model")
    p(doc, "Business and Membership define account scope. Category classifies Product. Each Product can have multiple Inventory Batches with quantity, receipt date, expiry date and storage location. Expiry Alert belongs to a batch. Stock Transaction records changes. Donation Log records completed handovers and deducts stock. Audit snapshots preserve useful history even if a batch is later deleted.")
    h(doc, "Controls that protect data quality")
    bullet(doc, "A stale version number blocks an edit that would overwrite a newer batch update.")
    bullet(doc, "Stock cannot become negative; CSV imports commit together or roll back together.")
    bullet(doc, "Business scope comes from the authenticated account, including for list and detail routes.")
    bullet(doc, "A POS retry reference avoids applying the same stock adjustment twice.")
    h(doc, "Suggested live walkthrough")
    for step in [
        "Sign in and review the dashboard counts and stock health.",
        "Create a product and batch, then filter the register by expiry.",
        "Import a CSV and show a rejected invalid row.",
        "Look up a batch by barcode or QR, adjust stock and inspect history.",
        "Review an alert and record a completed donation.",
    ]:
        doc.add_paragraph(step, "List Number")
    h(doc, "Known boundaries")
    p(doc, "The build uses SQLite and the Django development server for the local demo. A public rollout needs HTTPS, a production database, hosting, backups and deployment checks. Celery task logic is tested, but a live Redis worker and Beat have not been verified. No AI model, NGO marketplace or measured waste reduction is claimed in this milestone.")
    doc.save(OUT / "03_Project_Status_and_Technical_Walkthrough.docx")


if __name__ == "__main__":
    idea()
    mvp()
    status()
    for file in sorted(OUT.glob("*.docx")):
        print(file)
